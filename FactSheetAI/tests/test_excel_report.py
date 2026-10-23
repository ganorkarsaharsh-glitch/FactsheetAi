import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from factsheetai.excel_report import build, sample_config


class ExcelReportTemplateTests(unittest.TestCase):
    def test_reference_template_builds_formula_tabs_without_dynamic_array_functions(self):
        # sample_config is explicitly illustrative and is used only to exercise the template.
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as directory:
            path = Path(directory) / "template.xlsx"
            build(sample_config(), path)
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Summary", "Inputs", "Holdings", "Holdings_Prev", "Benchmark",
                    "Stock_Matrix", "Overlap", "Concentration", "Sector_Mcap", "Drift",
                    "Red_Flags", "Combined_Portfolio", "Validation", "Methodology",
                ],
            )
            formulas = [
                cell.value
                for sheet in workbook
                for row in sheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            self.assertGreater(len(formulas), 1000)
            joined = "\n".join(formulas).upper()
            for function in ("XLOOKUP(", "FILTER(", "SORT(", "UNIQUE("):
                self.assertNotIn(function, joined)
            self.assertTrue(workbook["Summary"]["A6"].value.startswith("="))
            workbook.close()


if __name__ == "__main__":
    unittest.main()
