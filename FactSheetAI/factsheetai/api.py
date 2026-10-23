"""FastAPI surface for source-backed portfolio comparison and availability."""

import json
import re
import shutil
import subprocess
import tempfile
from decimal import Decimal
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from starlette.background import BackgroundTask

from .amfi import load_top100
from .analytics import (
    assess_rule,
    combined_portfolio_detail,
    concentration_detail,
    equity_weights,
    health_score,
    holding_drift,
    identity_overlap,
    overlap,
    overlap_matrix,
    pandas_overlap,
    red_flags,
    sector_drift,
    style_drift,
    top_holdings_stats,
    turnover,
    validation_detail,
)
from .catalog import ROOT, SCHEMES, available_months, catalog_response, load_snapshot
from .excel_report import build as build_excel_report
from .report_metadata import extract_report_metadata

app = FastAPI(
    title="FactSheetAI",
    version="0.2.0",
    description="Deterministic, provenance-first portfolio analytics from AMFI and AMC disclosures",
)

STATIC_DIR = Path(__file__).parent / "static"
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

THRESHOLDS_PATH = ROOT / "config" / "analytics_thresholds.json"


def _thresholds() -> dict[str, Any]:
    return json.loads(THRESHOLDS_PATH.read_text(encoding="utf-8"))


def _prior_month(month: str) -> str:
    year, month_number = map(int, month.split("-"))
    return f"{year-1:04d}-12" if month_number == 1 else f"{year:04d}-{month_number-1:02d}"


def _top100_for_month(month: str) -> tuple[set[str], dict[str, str]]:
    config = json.loads((ROOT / "config" / "market_cap_lists.json").read_text(encoding="utf-8"))
    for record in config["lists"]:
        if month in record["covers_months"]:
            source = ROOT / record["path"]
            if not source.is_file():
                raise FileNotFoundError(f"Missing effective AMFI list: {record['path']}")
            listing = load_top100(source, record["effective_date"])
            source_url = _source_url(listing.source_file, "AMFI", None)
            if source_url is None:
                raise FileNotFoundError(f"Source URL missing from manifest: {listing.source_file}")
            return set(listing.isins), {
                "source_file": listing.source_file,
                "sha256": listing.sha256,
                "effective_date": listing.effective_date,
                "source_url": source_url,
            }
    raise FileNotFoundError(f"No effective AMFI half-yearly list configured for {month}")


def _source_url(source_file: str, amc: str, disclosure_month: str | None) -> str | None:
    manifest_path = ROOT / "data" / "raw" / "manifest.json"
    if not manifest_path.is_file():
        return None
    files = json.loads(manifest_path.read_text(encoding="utf-8"))["files"]
    for item in files:
        if Path(item["file"]).name != source_file or item["amc"] != amc:
            continue
        if disclosure_month is None or item.get("disclosure_month") == disclosure_month:
            return item["source_url"]
    return None


@app.get("/")
def home(view: str | None = None) -> FileResponse:
    if view in ("terminal", "app", "dashboard"):
        return FileResponse(Path(__file__).parent / "static" / "index.html")
    return FileResponse(Path(__file__).parent / "static" / "landing.html")


@app.get("/landing")
def landing() -> FileResponse:
    return FileResponse(Path(__file__).parent / "static" / "landing.html")


@app.get("/dashboard")
@app.get("/terminal")
@app.get("/app")
def dashboard() -> FileResponse:
    return FileResponse(Path(__file__).parent / "static" / "index.html")


@app.get("/logo.png")
def logo() -> FileResponse:
    return FileResponse(STATIC_DIR / "logo.png")


@app.get("/favicon.ico")
def favicon() -> FileResponse:
    return FileResponse(STATIC_DIR / "favicon.ico")


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/v1/catalog")
def catalog() -> dict[str, object]:
    return catalog_response()


@app.get("/v1/analysis")
def analyze(month: str, scheme_a: str, scheme_b: str) -> dict[str, object]:
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise HTTPException(status_code=422, detail="Month must use YYYY-MM format")
    try:
        a, b = load_snapshot(scheme_a, month), load_snapshot(scheme_b, month)
        result = overlap(a, b)
        top100, list_source = _top100_for_month(month)
    except (KeyError, FileNotFoundError, ValueError) as exc:
        raise HTTPException(status_code=404, detail=f"Data not available: {exc}") from exc

    source_a = a.holdings[0].provenance.model_dump(mode="json")
    source_b = b.holdings[0].provenance.model_dump(mode="json")
    source_a["source_url"] = _source_url(
        a.holdings[0].provenance.source_file, a.holdings[0].provenance.amc, month
    )
    source_b["source_url"] = _source_url(
        b.holdings[0].provenance.source_file, b.holdings[0].provenance.amc, month
    )
    overlap_sources = {"a": source_a, "b": source_b, "amfi_top100": list_source}
    common = []
    a_by_isin, b_by_isin = {h.isin: h for h in a.holdings}, {h.isin: h for h in b.holdings}
    for isin in result["common_isins"]:
        left, right = a_by_isin[isin], b_by_isin[isin]
        common.append(
            {
                "isin": isin,
                "name": left.security_name,
                "a_weight_pct": str(left.weight_pct),
                "b_weight_pct": str(right.weight_pct),
                "overlap_contribution_pct": str(min(left.weight_pct, right.weight_pct)),
                "source_rows": [left.provenance.source_row, right.provenance.source_row],
                "source_spans": [left.provenance.source_span, right.provenance.source_span],
            }
        )
    try:
        drift_a = style_drift(a, top100)
        drift_b = style_drift(b, top100)
        top_a, top_b = top_holdings_stats(a), top_holdings_stats(b)
    except (ValueError, ArithmeticError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    history = available_months(scheme_a)
    expected_prior = _prior_month(month)
    prior_month = expected_prior if expected_prior in history else None
    turnover_value: str | None = None
    if prior_month:
        try:
            turnover_value = str(turnover(load_snapshot(scheme_a, prior_month), a))
        except (FileNotFoundError, ValueError):
            pass
    return {
        "month": month,
        "schemes": [
            {"scheme_code": scheme_a, "name": a.scheme_name},
            {"scheme_code": scheme_b, "name": b.scheme_name},
        ],
        "overlap_pct": str(result["overlap_pct"]),
        "jaccard": str(result["jaccard"]),
        "common_holdings": common,
        "concentration": {
            scheme_a: {
                "top10_pct": str(top_a["top_n_concentration_pct"]),
                "hhi": str(top_a["equity_hhi"]),
                "effective_stocks": str(top_a["effective_stocks"]),
            },
            scheme_b: {
                "top10_pct": str(top_b["top_n_concentration_pct"]),
                "hhi": str(top_b["equity_hhi"]),
                "effective_stocks": str(top_b["effective_stocks"]),
            },
        },
        "drift": {
            scheme_a: {k: str(v) for k, v in drift_a.items()},
            scheme_b: {k: str(v) for k, v in drift_b.items()},
        },
        "active_share_pct": None,
        "active_share_note": "Unavailable: an authoritative benchmark portfolio was not supplied. Pairwise difference is not benchmark active share.",
        "turnover_a_pct": turnover_value,
        "red_flags": {
            scheme_a: red_flags(a, top100_isins=top100),
            scheme_b: red_flags(b, top100_isins=top100),
        },
        "validation": {
            "pandas_pct": str(pandas_overlap(a, b)),
            "pure_python_pct": str(result["overlap_pct"]),
            "identity_pct": str(identity_overlap(a, b)),
            "match": abs(pandas_overlap(a, b) - result["overlap_pct"]) <= Decimal("1e-9") and abs(identity_overlap(a, b) - result["overlap_pct"]) <= Decimal("1e-9"),
        },
        "provenance": overlap_sources,
    }


@app.get("/v1/thresholds")
def get_thresholds() -> dict[str, object]:
    return _thresholds()


@app.put("/v1/thresholds")
def update_thresholds(payload: dict[str, Any]) -> dict[str, object]:
    """Persist user-edited rule cutoffs; retain known rule structure and reject malformed values."""
    current = _thresholds()
    rules = payload.get("rules")
    if not isinstance(rules, dict) or set(rules) != set(current["rules"]):
        raise HTTPException(status_code=422, detail="Rules must include every configured rule code")
    for code, rule in rules.items():
        try:
            threshold = Decimal(str(rule["threshold"]))
            penalty = Decimal(str(rule["penalty"]))
            direction = rule["direction"]
        except (KeyError, ValueError, TypeError) as exc:
            raise HTTPException(status_code=422, detail=f"Invalid rule: {code}") from exc
        if not threshold.is_finite() or not penalty.is_finite() or penalty < 0 or direction not in {"above", "below"}:
            raise HTTPException(status_code=422, detail=f"Invalid threshold or penalty: {code}")
    margin = Decimal(str(payload.get("amber_margin_pct", 15)))
    if not margin.is_finite() or margin < 0 or margin > 100:
        raise HTTPException(status_code=422, detail="amber_margin_pct must be between 0 and 100")
    current["amber_margin_pct"] = float(margin)
    for code in rules:
        current["rules"][code]["threshold"] = float(rules[code]["threshold"])
        current["rules"][code]["penalty"] = float(rules[code]["penalty"])
        current["rules"][code]["direction"] = rules[code]["direction"]
        if code == "HIGH_PAIR_OVERLAP" and "moderate_threshold" in rules[code]:
            moderate = Decimal(str(rules[code]["moderate_threshold"]))
            if not moderate.is_finite() or moderate < 0 or moderate >= Decimal(str(rules[code]["threshold"])):
                raise HTTPException(status_code=422, detail="Moderate overlap threshold must be below the high threshold")
            current["rules"][code]["moderate_threshold"] = float(moderate)
    THRESHOLDS_PATH.write_text(json.dumps(current, indent=2) + "\n", encoding="utf-8")
    return current


@app.post("/v1/upload/preview")
def upload_preview(payload: dict[str, str]) -> dict[str, object]:
    """Parse only the project's normalized CSV format and return a human review record."""
    csv_text = payload.get("csv")
    if not csv_text or len(csv_text.encode("utf-8")) > 12_000_000:
        raise HTTPException(status_code=413, detail="Provide a normalized CSV up to 12 MB")
    from .ingestion import load_csv

    upload_dir = Path(tempfile.mkdtemp(prefix=".upload-preview-", dir=ROOT))
    temp_path = upload_dir / "review.csv"
    try:
        with temp_path.open("w", encoding="utf-8", newline="") as f:
            f.write(csv_text)
        snapshot = load_csv(temp_path)
        total = sum((h.weight_pct for h in snapshot.holdings), Decimal(0))
    except (ValueError, OSError) as exc:
        raise HTTPException(status_code=422, detail=f"CSV could not be reviewed: {exc}") from exc
    finally:
        shutil.rmtree(upload_dir, ignore_errors=True)
    return {"status": "REVIEW_REQUIRED", "scheme_id": snapshot.scheme_id,
            "scheme_name": snapshot.scheme_name, "disclosure_month": snapshot.disclosure_month.isoformat(),
            "holdings_count": len(snapshot.holdings), "weights_pct": str(total),
            "weights_near_100": abs(total - Decimal(100)) <= Decimal("0.5"),
            "source_file": Path(payload.get("filename", "upload.csv")).name,
            "sha256": snapshot.holdings[0].provenance.sha256,
            "parsed_at": snapshot.holdings[0].provenance.parsed_at.isoformat(),
            "extraction_method": "normalized CSV parser", "confidence": "source row parsed; human confirmation required",
            "sample_holdings": [{"isin": h.isin, "name": h.security_name, "weight_pct": str(h.weight_pct),
                                 "source_row": h.provenance.source_row} for h in snapshot.holdings[:20]],
            "limitation": "PDF factsheets and arbitrary AMC Excel layouts require adapter-specific extraction and review. This endpoint does not import the preview into the official corpus."}


@app.get("/v1/report-metadata")
def report_metadata(month: str, schemes: str, research: bool = False) -> dict[str, Any]:
    """Return automatically extracted, source-grounded metadata for the selected report."""
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise HTTPException(status_code=422, detail="Month must use YYYY-MM format")
    codes = list(dict.fromkeys(code.strip() for code in schemes.split(",") if code.strip()))
    if not 2 <= len(codes) <= 4:
        raise HTTPException(status_code=422, detail="Select between two and four distinct schemes")
    try:
        return {
            "month": month,
            "schemes": extract_report_metadata(month, codes, research_missing=research),
        }
    except (KeyError, FileNotFoundError, ValueError) as exc:
        raise HTTPException(status_code=404, detail=f"Metadata source unavailable: {exc}") from exc


@app.post("/v1/export/excel")
def export_excel(payload: dict[str, Any]) -> FileResponse:
    """Build an Excel report from source holdings and metadata extracted from their AMC workbooks."""
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice and not payload.get("allow_uncalculated", False):
        raise HTTPException(
            status_code=503,
            detail="Excel export requires a headless LibreOffice runtime to recalculate and validate formula outputs.",
        )
    month = str(payload.get("month", ""))
    raw_codes = payload.get("schemes", [])
    if not isinstance(raw_codes, list) or any(not isinstance(code, str) for code in raw_codes):
        raise HTTPException(status_code=422, detail="schemes must be a list of scheme codes")
    codes = list(dict.fromkeys(raw_codes))
    investments = payload.get("investments", {})
    research_missing = payload.get("research_missing", False)
    if not isinstance(research_missing, bool):
        raise HTTPException(status_code=422, detail="research_missing must be a boolean")
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month) or not 2 <= len(codes) <= 4:
        raise HTTPException(status_code=422, detail="Provide a valid month and two to four distinct schemes")
    if not isinstance(investments, dict):
        raise HTTPException(status_code=422, detail="investments must be a JSON object")
    try:
        metadata = extract_report_metadata(month, codes, research_missing=research_missing)
    except (KeyError, FileNotFoundError, ValueError) as exc:
        raise HTTPException(status_code=404, detail=f"Metadata source unavailable: {exc}") from exc
    for code in codes:
        try:
            amount = Decimal(str(investments.get(code, 0)))
        except (ValueError, TypeError) as exc:
            raise HTTPException(status_code=422, detail=f"Investment amount is invalid for {code}") from exc
        if not amount.is_finite() or amount < 0:
            raise HTTPException(status_code=422, detail=f"Investment amount must be finite and non-negative for {code}")
    benchmark_code = str(payload.get("benchmark_scheme", "119063"))
    if benchmark_code not in SCHEMES or SCHEMES[benchmark_code].get("role") != "benchmark_control":
        raise HTTPException(status_code=422, detail="Select a configured benchmark-control scheme")
    try:
        snapshots = {code: load_snapshot(code, month) for code in codes}
        benchmark = load_snapshot(benchmark_code, month)
        top100, cap_source = _top100_for_month(month)
    except (KeyError, FileNotFoundError, ValueError) as exc:
        raise HTTPException(status_code=404, detail=f"Data not available for report: {exc}") from exc
    prior = _prior_month(month)
    try:
        previous = {code: load_snapshot(code, prior) for code in codes}
    except (FileNotFoundError, ValueError) as exc:
        raise HTTPException(status_code=404, detail=f"Immediate prior-month holdings are required for the workbook: {exc}") from exc
    canonical_names: dict[str, str] = {}
    canonical_sectors: dict[str, str] = {}
    for snapshot in [*snapshots.values(), *previous.values(), benchmark]:
        for holding in snapshot.holdings:
            if holding.asset_type != "equity":
                continue
            canonical_names.setdefault(holding.isin, holding.security_name)
            if holding.sector:
                canonical_sectors.setdefault(holding.isin, holding.sector)

    def workbook_rows(snapshot: Any) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        for holding in snapshot.holdings:
            if holding.asset_type != "equity":
                continue
            rows.append({"stock": f"{canonical_names[holding.isin]} [{holding.isin}]", "sector": canonical_sectors.get(holding.isin, "Unclassified"),
                         "mcap": "Large" if holding.isin in top100 else "Unknown",
                         "weight": float(holding.weight_pct / Decimal(100))})
        return rows

    funds: list[dict[str, Any]] = []
    lineage: list[dict[str, Any]] = []
    for index, code in enumerate(codes):
        info = metadata[code]
        financials: dict[str, float | None] = {}
        for key in ("aum_cr", "expense", "cat_expense", "turnover", "cash"):
            raw_value = info.get(key)
            if raw_value is None:
                financials[key] = None
                continue
            try:
                number = float(raw_value)
            except (ValueError, TypeError) as exc:
                raise HTTPException(status_code=422, detail=f"Numeric metadata is invalid for {code}: {key}") from exc
            if not Decimal(str(number)).is_finite() or (key != "cash" and number < 0):
                raise HTTPException(status_code=422, detail=f"Numeric metadata is outside its valid range for {code}: {key}")
            financials[key] = number
        funds.append({"short": f"Fund {chr(65+index)}", "name": snapshots[code].scheme_name,
                      "category": str(info["category"]), "mandate_large_min": float(json.loads((ROOT / "config" / "mandates.json").read_text(encoding="utf-8"))["large_cap_floor_pct"]) / 100,
                      "aum_cr": financials["aum_cr"], "expense": financials["expense"] / 100 if financials["expense"] is not None else None,
                      "cat_expense": financials["cat_expense"] / 100 if financials["cat_expense"] is not None else None,
                      "turnover": financials["turnover"] / 100 if financials["turnover"] is not None else None,
                      "cash": financials["cash"] / 100 if financials["cash"] is not None else None,
                      "investment": float(investments.get(code, 0)),
                      "manager": info.get("manager") or "Unavailable", "holdings": workbook_rows(snapshots[code]),
                      "prev_holdings": workbook_rows(previous[code])})
        lineage.append({"item": snapshots[code].scheme_name, "source": str(info["source"]),
                        "page_table": str(info["page_table"]), "extracted_on": month,
                        "method": "Automatic extraction from the selected AMC monthly portfolio workbook",
                        "confidence": "source-parsed; fields absent from source marked unavailable"})
        for field, citation in info.get("field_sources", {}).items():
            lineage.append({
                "item": f"{snapshots[code].scheme_name} — {field}",
                "source": citation["url"],
                "page_table": f"{citation['title']} — Evidence: {citation['evidence']}",
                "extracted_on": month,
                "method": "OpenAI Responses web search, restricted to official AMC domain; source URL and value retained",
                "confidence": "official-domain citation validated; exact-month value extracted",
            })
    if len({row["stock"] for fund in funds for row in fund["holdings"]}) < 10:
        raise HTTPException(status_code=422, detail="Workbook requires at least ten distinct current holdings")
    bench_name = f"{benchmark.scheme_name} (disclosed index-fund proxy)"
    lineage.append({"item": bench_name, "source": benchmark.holdings[0].provenance.source_file,
                    "page_table": benchmark.holdings[0].provenance.source_span or "source rows",
                    "extracted_on": benchmark.holdings[0].provenance.parsed_at.date().isoformat(),
                    "method": "AMC adapter; proxy, not official index weights", "confidence": "source-parsed"})
    report = {"report_title": "FactSheetAI - Overlap & Drift Report", "as_of": month,
              "prior_month": prior, "benchmark_name": bench_name,
              "source_note": f"AMFI top-100: {cap_source['source_file']}; benchmark is an index-fund proxy.",
              "data_status": "Source-backed holdings and automatically extracted metadata; values absent from the workbook were searched in official AMC factsheets when requested and remain unavailable without a verified citation; outside AMFI top-100 is unclassified; index fund is a benchmark proxy.",
              "funds": funds, "benchmark": workbook_rows(benchmark), "lineage": lineage}
    temp_dir = tempfile.mkdtemp(prefix=".factsheetai-report-", dir=ROOT)
    input_path = Path(temp_dir) / "report_input.xlsx"
    output_path = Path(temp_dir) / "FactSheetAI_Report.xlsx"
    try:
        build_excel_report(report, input_path)
        from openpyxl import load_workbook  # type: ignore[import-untyped]
        if libreoffice:
            recalc_dir = Path(temp_dir) / "recalculated"
            recalc_dir.mkdir()
            completed = subprocess.run(
                [libreoffice, "--headless", "--convert-to", "xlsx", "--outdir", str(recalc_dir), str(input_path)],
                capture_output=True, text=True, timeout=90, check=False,
            )
            recalculated = recalc_dir / input_path.name
            if completed.returncode != 0 or not recalculated.is_file():
                raise RuntimeError(completed.stderr or completed.stdout or "LibreOffice recalculation failed")
            checked = load_workbook(recalculated, data_only=True, read_only=True)
            errors = [
                (sheet.title, cell.coordinate, cell.value)
                for sheet in checked
                for row in sheet.iter_rows()
                for cell in row
                if cell.data_type == "e"
                or (isinstance(cell.value, str) and cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")))
            ]
            checked.close()
            if errors:
                raise RuntimeError(f"Formula errors found: {errors[:10]}")
            shutil.copy2(recalculated, output_path)
        else:
            # Excel and other spreadsheet apps recalculate formulas when the file opens.
            # Validate the generated package and preserve formulas without requiring LibreOffice.
            shutil.copy2(input_path, output_path)
            checked = load_workbook(output_path, data_only=False, read_only=True)
            if "Summary" not in checked.sheetnames or "Holdings" not in checked.sheetnames:
                raise RuntimeError("Generated workbook is missing required report sheets")
            checked.close()
    except (OSError, ValueError, RuntimeError, subprocess.TimeoutExpired) as exc:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise HTTPException(status_code=422, detail=f"Excel report validation failed: {exc}") from exc
    return FileResponse(output_path, filename="FactSheetAI_Report.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        background=BackgroundTask(shutil.rmtree, temp_dir, ignore_errors=True))


@app.get("/v1/portfolio-analysis")
def analyze_portfolio(month: str, schemes: str, investments: str | None = None) -> dict[str, object]:
    """Analyze two to four selected disclosures with deterministic, source-backed metrics."""
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise HTTPException(status_code=422, detail="Month must use YYYY-MM format")
    codes = list(dict.fromkeys(x.strip() for x in schemes.split(",") if x.strip()))
    if not 2 <= len(codes) <= 4:
        raise HTTPException(status_code=422, detail="Select between two and four distinct schemes")
    try:
        snapshots = {code: load_snapshot(code, month) for code in codes}
        top100, source = _top100_for_month(month)
    except (KeyError, FileNotFoundError, ValueError) as exc:
        raise HTTPException(status_code=404, detail=f"Data not available: {exc}") from exc
    matrix = overlap_matrix(snapshots)
    thresholds = _thresholds()
    high_overlap = Decimal(str(thresholds["rules"]["HIGH_PAIR_OVERLAP"]["threshold"]))
    moderate_overlap = Decimal(str(thresholds["rules"]["HIGH_PAIR_OVERLAP"].get("moderate_threshold", 30)))
    pairs = []
    for i, left in enumerate(codes):
        for right in codes[i + 1:]:
            value = matrix[left][right]
            left_isins, right_isins = set(equity_weights(snapshots[left])), set(equity_weights(snapshots[right]))
            common_isins = left_isins & right_isins
            pairs.append({"scheme_a": left, "scheme_b": right, "overlap_pct": str(value),
                          "common_count": len(common_isins), "unique_a_count": len(left_isins-common_isins),
                          "unique_b_count": len(right_isins-common_isins),
                          "unmatched_a_isins": sorted(left_isins-common_isins),
                          "unmatched_b_isins": sorted(right_isins-common_isins),
                          "rank": "HIGH" if value >= high_overlap else "MODERATE" if value >= moderate_overlap else "LOW"})
    alloc = {}
    if investments:
        try:
            raw = json.loads(investments)
            alloc = {key: Decimal(str(value)) for key, value in raw.items()}
        except (ValueError, TypeError) as exc:
            raise HTTPException(status_code=422, detail="investments must be a JSON object of rupee amounts") from exc
    if not alloc:
        alloc = {code: Decimal(1) for code in codes}
    try:
        combined = combined_portfolio_detail(snapshots, alloc)
        details = {code: concentration_detail(snapshots[code]) for code in codes}
        style = {code: style_drift(snapshots[code], top100) for code in codes}
        prev_months = {
            code: (_prior_month(month) if _prior_month(month) in available_months(code) else None)
            for code in codes
        }
        monthly = {}
        for code, prior in prev_months.items():
            if prior:
                previous = load_snapshot(code, prior)
                monthly[code] = {**holding_drift(previous, snapshots[code]),
                                 "sector_changes": sector_drift(previous, snapshots[code])}
        validation = validation_detail(snapshots)
    except (ValueError, ArithmeticError, FileNotFoundError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    rule_defs = thresholds["rules"]
    flags_by_fund: dict[str, list[dict[str, object]]] = {}
    health_by_fund: dict[str, dict[str, object]] = {}
    floor = Decimal(str(json.loads((ROOT / "config" / "mandates.json").read_text(encoding="utf-8"))["large_cap_floor_pct"]))
    for code in codes:
        candidates: dict[str, Decimal | None] = {
            "HIGH_PAIR_OVERLAP": max((Decimal(str(p["overlap_pct"])) for p in pairs if code in {p["scheme_a"], p["scheme_b"]}), default=None),
            "MANDATE_BREACH": Decimal(str(style[code]["top100_weight_pct"])),
            "SINGLE_STOCK_CONCENTRATION": Decimal(str(details[code]["largest_holding"])),
            "TOP10_CONCENTRATION": Decimal(str(details[code]["top10_pct"])),
            "SECTOR_CROWDING": (Decimal(str(details[code]["largest_sector_pct"]))
                                 if Decimal(str(details[code]["sector_coverage_pct"])) >= Decimal(100)
                                 else None),
            "FEW_EFFECTIVE_STOCKS": Decimal(str(details[code]["effective_stocks"])),
            "HEAVY_RESHUFFLING": Decimal(str(monthly[code]["portfolio_change_pct"])) if code in monthly else None,
            "CLOSET_INDEXING": None, "EXPENSE_ABOVE_CATEGORY": None, "CASH_DRAG": None,
            "HIGH_TURNOVER": None,
        }
        fund_flags = []
        for rule_code, metric_value in candidates.items():
            definition = dict(rule_defs[rule_code])
            if rule_code == "MANDATE_BREACH":
                definition["threshold"] = floor
            severity = assess_rule({**definition, "amber_margin_pct": thresholds["amber_margin_pct"]}, metric_value)
            fund_flags.append({"code": rule_code, "rule": definition["label"],
                               "threshold": str(definition["threshold"]),
                               "actual": str(metric_value) if metric_value is not None else None,
                               "severity": severity, "unit": definition["unit"]})
        flags_by_fund[code] = fund_flags
        health_by_fund[code] = health_score(fund_flags, rule_defs)
    pair_flags = [{"code": "HIGH_PAIR_OVERLAP", "rule": rule_defs["HIGH_PAIR_OVERLAP"]["label"],
                   "threshold": str(rule_defs["HIGH_PAIR_OVERLAP"]["threshold"]),
                   "actual": pair["overlap_pct"],
                    "severity": assess_rule({**rule_defs["HIGH_PAIR_OVERLAP"], "amber_margin_pct": thresholds["amber_margin_pct"]}, Decimal(str(pair["overlap_pct"]))),
                   "scheme_a": pair["scheme_a"], "scheme_b": pair["scheme_b"]} for pair in pairs]
    alternatives: dict[str, Any] = {}
    for candidate_code in SCHEMES:
        if candidate_code in codes or month not in available_months(candidate_code):
            continue
        try:
            alternatives[candidate_code] = load_snapshot(candidate_code, month)
        except (FileNotFoundError, ValueError):
            continue
    simulations = []
    if alternatives:
        for dropped in codes:
            rest = [code for code in codes if code != dropped]
            before = sum((overlap(snapshots[dropped], snapshots[other])["overlap_pct"] for other in rest), Decimal(0))
            for candidate_code, candidate in alternatives.items():
                after = sum((overlap(candidate, snapshots[other])["overlap_pct"] for other in rest), Decimal(0))
                simulations.append({"replace": dropped, "with": candidate_code, "pairwise_overlap_before_pct": str(before),
                                    "pairwise_overlap_after_pct": str(after), "change_pct": str(after-before),
                                    "label": "Analytical simulation only; not investment advice."})
    simulations.sort(key=lambda item: (Decimal(item["change_pct"]), item["replace"], item["with"]))
    def jsonable(value: Any) -> Any:
        if isinstance(value, Decimal):
            return str(value)
        if isinstance(value, dict):
            return {key: jsonable(item) for key, item in value.items()}
        if isinstance(value, list):
            return [jsonable(item) for item in value]
        return value
    return {"month": month, "schemes": [{"scheme_code": code, "name": snapshots[code].scheme_name} for code in codes],
            "overlap_matrix": {a: {b: str(v) for b, v in row.items()} for a, row in matrix.items()},
            "pairs": pairs, "concentration": {k: {x: str(v) if isinstance(v, Decimal) else v for x, v in val.items()} for k, val in details.items()},
            "style_drift": {k: {x: str(v) if isinstance(v, Decimal) else v for x, v in val.items()} for k, val in style.items()},
            "market_cap_split": {k: {"large_top100_pct": str(style[k]["top100_weight_pct"]),
                                     "outside_top100_unclassified_pct": str(style[k]["outside_top100_weight_pct"]),
                                     "mid_pct": None, "small_pct": None,
                                     "limitation": "AMFI top-100 evidence identifies large-cap exposure; outside-top-100 holdings cannot be split into mid/small from this corpus."} for k in codes},
            "monthly_drift": {k: {x: [{q: str(v) if isinstance(v, Decimal) else v for q, v in row.items()} for row in value] if isinstance(value, list) else str(value) if isinstance(value, Decimal) else value for x, value in drift.items()} for k, drift in monthly.items()},
            "combined": jsonable(combined),
            "expense": None, "expense_note": "Unavailable: expense disclosures are not included in the normalized source records.",
            "active_share": None, "active_share_note": "Unavailable: no authoritative benchmark disclosure is loaded.",
            "red_flags": flags_by_fund, "health_scores": {k: {x: str(v) for x, v in score.items()} for k, score in health_by_fund.items()},
            "pair_flags": pair_flags,
            "swap_simulations": simulations[:10],
            "validation": {**validation, "pair_recomputations": {f"{a}/{b}": {"min_sum": str(overlap(snapshots[a], snapshots[b])["overlap_pct"]), "identity_sum": str(identity_overlap(snapshots[a], snapshots[b])), "match": identity_overlap(snapshots[a], snapshots[b]) == overlap(snapshots[a], snapshots[b])["overlap_pct"]} for i, a in enumerate(codes) for b in codes[i+1:]}},
            "benchmark_source": None, "amfi_top100_source": source,
            "thresholds": thresholds, "source_status": "Authoritative monthly portfolio snapshots; cash and benchmark analytics unavailable unless disclosed."}
